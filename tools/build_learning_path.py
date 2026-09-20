"""Build the curated Fall 2026 release from user-supplied teaching sources.
Run with --input pointing to the extracted Fall2026 directory.
The source ledger preserves every supplied slide, including title/summary pages.
"""
from pathlib import Path
import argparse,base64,copy,hashlib,html,importlib.util,json,re,shutil,subprocess,tempfile
from bs4 import BeautifulSoup
from pptx import Presentation
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('pathspec',ROOT/'curriculum/learning-path/spec.py');mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
CHAPTERS,ADDITIONS=mod.CHAPTERS,mod.ADDITIONS
study_spec=importlib.util.spec_from_file_location('studyspec',ROOT/'curriculum/learning-path/study.py');study_mod=importlib.util.module_from_spec(study_spec);study_spec.loader.exec_module(study_mod)
BUILD='20260921-mastery-v2'
def clean(x):return BeautifulSoup(str(x),'html.parser').get_text(' ',strip=True)
def esc(x):return html.escape(str(x),quote=True)
def textwalk(x):
 out=[]
 if isinstance(x,dict):
  for k,v in x.items():
   if k=='txt' and isinstance(v,str):out.append(clean(v))
   elif isinstance(v,(dict,list)):out+=textwalk(v)
 elif isinstance(x,list):
  for v in x:out+=textwalk(v)
 return out

def source_slides(p):
 if p.suffix=='.pdf':
  r=PdfReader(p);return [{'slide':i+1,'title':(pg.extract_text() or '').split('\n')[0],'text':pg.extract_text() or '', 'tables':[],'images':[]} for i,pg in enumerate(r.pages)]
 out=[]
 def shapes(items,tx,tab,imgs):
  for sh in items:
   if sh.shape_type==6:shapes(sh.shapes,tx,tab,imgs)
   if sh.has_text_frame:tx.append(sh.text)
   if sh.has_table:tab.append([[c.text for c in row.cells] for row in sh.table.rows])
   if sh.shape_type==13:
    im=sh.image
    if im.ext.lower() in ['png','jpg','jpeg','gif','webp']:imgs.append('data:'+im.content_type+';base64,'+base64.b64encode(im.blob).decode())
 for i,s in enumerate(Presentation(p).slides):
  tx=[];tabs=[];imgs=[];shapes(s.shapes,tx,tabs,imgs)
  title=next((t for t in tx if t.strip() and not re.fullmatch(r'\d+|\d+/\d+/\d+',t.strip())),'Source figure')
  out.append({'slide':i+1,'title':title,'text':'\n'.join(tx),'tables':tabs,'images':imgs})
 return out

def rawunit(k,title,body,route='core',**more):
 return dict(k=k,title=title,phase='LEARN',route=route,blocks=[dict(t='raw',html=body)],noAnswer=True,**more)
def source_html(sl):
 t=sl['text'];lines=[x for x in t.splitlines() if not re.fullmatch(r'\s*\d+\s*|\d{1,2}/\d{1,2}/\d{2,4}',x.strip()) and not re.match(r'^Chapter \d+ ',x)]
 body='<div class="lp-source-text">'+esc('\n'.join(lines))+'</div>'
 for table in sl['tables']:
  body+='<div class="lp-tablewrap"><table>'+''.join('<tr>'+''.join('<td>'+esc(c).replace('\n','<br>')+'</td>' for c in row)+'</tr>' for row in table)+'</table></div>'
 for im in sl['images']:body+='<img loading="lazy" src="'+im+'" alt="Original embedded figure from supplied source slide '+str(sl['slide'])+'">'
 return '<details class="lp-source-section" id="source-slide-'+str(sl['slide'])+'"><summary>Slide '+str(sl['slide'])+' · '+esc(sl['title'])+'</summary>'+body+'</details>'

def addunit(k):
 title,items,a,b,q,ans=ADDITIONS[k]
 return dict(k=k,title=title,sub='Required source concepts with a short worked explanation.',blocks=[dict(t='cards',cols=3,items=[dict(c='teal',lab=lab,txt=txt,src=f'P1 · pp.{a}–{b}') for lab,txt in items])],task=q,answer=ans,sourceRange=f'{a}–{b}',noAnswer=False)

def repair_figures(n,L):
 maps={11:{'X01':'ch11_costcurve','X02':'ch11_iomap','X03':'ch11_usage','R09':'ch11_selfmon','R08':'ch11_airbus','X09':'ch11_relmeasure','X10':'ch11_opprofile'},12:{'X01':'ch12_risktri','X04':'ch12_risktri','X05':'ch12_faulttree','R18':'ch12_modelcheck','R08':'ch12_structarg','R10':'ch12_hierarchy','X10':'ch12_safetyarg'},13:{'X01':'ch13_layers','X02A':'ch13_layers','X03':'ch13_riskproc','X06':'ch13_misuse','R11':'ch13_designrisk','R09':'ch13_layered','B01':'ch13_layered'},14:{'X01':'ch14_resactivities','X02':'ch14_resactivities','X02A':'ch14_cyberplan','X03':'ch14_nested','X04':'ch14_swiss','X06':'ch14_survivability','R18':'ch14_clientserver','B01':'ch14_mentcare'},15:{'X01':'ch15_landscape','X03':'ch15_mvc','X04':'ch15_basesystem','X05':'ch15_instance','R18':'ch15_erparch','R09':'ch15_wrapping'},16:{'X01':'ch16_interfaces','X02A':'ch16_datacollector','X03':'ch16_interfaces','X04':'ch16_modelelements','X05':'ch16_cbseprocess','X06':'ch16_withreuse','R18':'ch16_identification','R09':'ch16_adaptor','B01':'ch16_composition'},17:{'X02':'ch17_middleware','X02A':'ch17_layered','X03':'ch17_middleware','X04':'ch17_middleware','X05':'ch17_layered','X06':'ch17_masterslave','X07':'ch17_thinfat','X09':'ch17_saas'},20:{'X01':'ch20_complexity','X03':'ch20_reality','X04':'ch20_soseng','X05':'ch20_serviceiface','X06':'ch20_ilearn','X09':'ch20_togaf'}}
 # Only change figure bindings where the supplied clone had mismatched content.
 captions={'ch16_interfaces':'slide 15 · provides and requires interfaces','ch16_datacollector':'slide 16 · data collector component','ch16_modelelements':'slide 19 · component-model elements','ch16_cbseprocess':'slide 25 · CBSE processes','ch16_withreuse':'slide 35 · development with reuse','ch16_adaptor':'slide 49 · adapter connecting collector and sensor','ch16_composition':'slide 43 · composition types'}
 if n==16:
  for u in L['units']:
   figs=[b for b in u.get('blocks',[]) if b.get('t')=='figure']
   if u['k'] in maps[n] and figs:
    f=figs[0];key=maps[n][u['k']];f.update(img=key,cap=captions.get(key,'supplied source diagram'),right=[]);f.pop('native',None)
    if u['k']=='X02A':u['blocks']=[b for b in u['blocks'] if b.get('t')!='figure']+[f]
 # Correct cloned figure placement in the distributed and SoS chapters.
 adjustments={
  17:{'X04':('ch17_middleware','slide 21 · middleware'),'X05':('ch17_layered','slide 27 · client-server layers'),'X06':('ch17_masterslave','slide 32 · traffic management architecture')},
  20:{'X03':('ch20_reality','slide 25 · system-of-systems reality'),'X04':('ch20_soseng','slide 29 · SoS engineering'),'X05':('ch20_serviceiface','slide 34 · service interfaces'),'X06':('ch20_ilearn','slide 38 · staged iLearn deployment')}
 }
 for u in L['units']:
  for b in u.get('blocks',[]):
   if b.get('t')=='figure' and b.get('native')=='props' and n!=10:
    b['srcName']='Chapter 10 recap · course redraw';b['cap']='Five dependability properties; not a figure from this chapter'
  if u['k'] in adjustments.get(n,{}):
   figs=[b for b in u.get('blocks',[]) if b.get('t')=='figure']
   if figs:
    key,caption=adjustments[n][u['k']];figs[0].update(img=key,cap=caption,right=[]);figs[0].pop('native',None)
  if n==17 and u['k']=='X03':
   u['blocks']=[b for b in u['blocks'] if b.get('t')!='figure']
   u['blocks'].insert(0,dict(t='conceptflow',items=[['Caller','Sends an operation request'],['Remote service','May execute independently'],['Response','Can arrive, fail or time out']]))
  if n==20 and u['k']=='X09':u['blocks']=[b for b in u['blocks'] if b.get('t')!='figure']
  if n==20 and u['k']=='X10':u['blocks'].insert(0,dict(t='figure',img='ch20_togaf',cap='slide 46 · TOGAF Architecture Development Method',right=[]))
  if n==20 and u['k']=='X02A':u['blocks'].insert(0,dict(t='figure',img='ch20_classification',cap='slide 18 · SoS classification',right=[]))
  if n==17 and u['k']=='X10':u['blocks'].append(dict(t='wide',c='sand',lab='DISTINCT DESIGN QUESTIONS',txt='SaaS describes how software is delivered. SOA describes how a system is structured. A SaaS application can use a service-oriented architecture. The source compares typical interaction patterns; statefulness alone does not define either category.'))
 # Source-era technology examples are kept, but categorical modern claims removed.
 def fix(x):
  if isinstance(x,list):return[fix(v) for v in x]
  if isinstance(x,dict):return{k:fix(v) for k,v in x.items()}
  if isinstance(x,str):
   x=x.replace('Multiple competing standards killed CBSE’s promise of universal reuse. Service-oriented SE is replacing it.','The source describes competing component standards and service-oriented alternatives. Compare their integration assumptions; these historical examples do not establish current market adoption.')
   x=x.replace('Services are based around standards, so there are no communication problems.','Services use shared communication standards, but interoperability, semantics and failures still require validation.')
   x=x.replace('It is impossible for components developed using different approaches to work together.','Components from different models may require bridging, wrappers or adapters rather than direct composition.')
   return x
  return x
 L.update(fix(L))

def build(base):
 catalog=[];audit=[]
 css=(ROOT/'curriculum/learning-path/reader.css').read_text()+'\n'+(ROOT/'curriculum/learning-path/study.css').read_text();runtime=(ROOT/'curriculum/learning-path/reader.js').read_text()+'\n'+(ROOT/'curriculum/learning-path/study.js').read_text()
 for ai,(n,c) in enumerate(CHAPTERS.items(),1):
  source=next(p for p in base.glob(f'Ch{n} *') if p.suffix in ['.pptx','.pdf']);src_html=next(base.glob(f'Ch{n}-*.html'));markup=src_html.read_text()
  soup=BeautifulSoup(markup,'html.parser');script=soup.find_all('script')[0].string
  part=script.rfind('/*',0,script.index('PART 2 — ENGINE'));prefix=script[:script.index('var LECTURE=')];engine=script[script.index('var V=LECTURE',part):]
  # The pre-existing versioned route mutators are replaced by the curated route.
  program='const vm=require("vm"),fs=require("fs");const c={};vm.createContext(c);vm.runInContext(fs.readFileSync(process.argv[1],"utf8"),c);process.stdout.write(JSON.stringify(c.LECTURE));'
  with tempfile.NamedTemporaryFile(mode='w',suffix='.js') as tmp:
   tmp.write(script[:part]);tmp.flush();L=json.loads(subprocess.check_output(['node','-e',program,tmp.name],text=True))
  repair_figures(n,L);by={u['k']:u for u in L['units']};slides=source_slides(source);count=len(slides)
  titles=json.loads((ROOT/'curriculum/learning-path/titles.json').read_text()).get(str(n),{})
  for k,t in titles.items():
   if k in by:by[k]['title']=t
  destsource=f'lectures/iscarb/sources/Ch{n}-Original{source.suffix}';shutil.copy2(source,ROOT/destsource)
  L['path']=dict(chapter=n,assignment=ai,exit=c['exit'],sourceFile='sources/Ch'+str(n)+'-Original'+source.suffix,build=BUILD)
  L['study']=study_mod.STUDY[n]
  L['course']=f'CPIT-455 / CHAPTER {n}';L['title']=c['title'];L['titleLead']=c['title'];L['titleAccent']=''
  L['tagline']=c['challenge']+' Learn the source concepts, try three short interactions, then apply them in a separate assignment.'
  L['sourceMap']=f'Supplied original Chapter {n} · {count} source slides · complete source ledger in Source review'
  L['phases']=['LEARN','APPLY','CHECK'];L['flow']='LEARN → PRACTISE → REVIEW → TRANSFER';L['cardSub']='Reusable practice artifact · six fields · save or export before leaving'
  L['context']['schedule']='See the current LMS timetable.';L['context']['assignmentDue']='the deadline announced in the LMS';L['context']['contact']='arnoor@kau.edu.sa';L['context']['readingSource']=f'Supplied Chapter {n} source pack; every slide is indexed under Source review.'
  L['context']['bodies']='Identify the relevant authority for the actual deployment context.';L['context']['bodiesShort']='Applicable sector and institutional authorities'
  L['preparation']=dict(source=L['context']['readingSource'],due='before the assignment deadline in the LMS',ledger='C01')
  # Five chapter objectives are local teaching objectives, not replacements for the eleven syllabus CLOs.
  for r in L.get('readiness',[])[:5]:r['clo']=r.get('clo','').replace('CLO','LO');r['plo']='Practice';r['where']='Classroom practice and source review'
  L['readiness']=L.get('readiness',[])[:5]
  L['coverage']={'source':f'Supplied Sommerville Chapter {n} slides','span':[1,count],'sections':[dict(k=str(i+1),t=g[0],pp=[g[1],min(g[2],count)],concept=g[0]) for i,g in enumerate(c['groups'])]}
  for u in by.values():
   u['route']='toolkit';u['appendix']=True;u['phase']='ISCARB TOOLKIT';u.pop('simpleMain',None);u.pop('simpleKind',None)
   if u.get('rule'):u['rule']=u['rule'].replace('CLO','LO')
  # Chapter-only practice, without unverified official mappings or exam-weight claims.
  by['J01']=rawunit('J01','Learning outcomes and readiness alignment','<div class="lp-callout"><p><b>Five chapter objectives</b> support syllabus CLO '+', '.join(map(str,c['clo']))+'. Knowledge is checked through explanation; skills through the chapter artifact; responsibility through evidence limits, AI disclosure and human review.</p><p>NCAAA / Jaheziah: this is a teaching alignment view. The supplied files do not establish an approved current national item-weight or departmental mapping. Course-written practice questions are not official exam items or a prediction of national-test performance.</p></div>',route='toolkit')
  by['J02']['title']='Chapter practice questions';by['J02']['sub']='Try first, then read the explanation. These are course-written practice items, not official national-test questions.';by['J02']['task']='Answer independently and revisit the relevant source topic for each error.';by['J02']['notes']='Use for individual review after class. Do not infer national test performance from this score.'
  by['R20']['sub']='The gates are self-assessed. Filled fields indicate that a response exists; the instructor must judge whether its reasoning is sound.';by['R20']['notes']='No automatic check certifies mastery. Review the chapter outcomes and the actual quality of evidence.'
  for key in c['core']:
   if key in ADDITIONS:by[key]=addunit(key)
   u=by[key];u.update(route='core',appendix=False,phase='LEARN',rule='',showIf=None)
   hits=[g for g in c['groups'] if key in g[3]]
   if not u.get('sourceRange') and hits:u['sourceRange']='; '.join(f'{g[1]}–{min(g[2],count)}' for g in hits)
   if key not in ADDITIONS:
    statements=textwalk(u.get('blocks',[]));technical=[t for t in statements if len(t)>35]
    if not u.get('answer'):
     # A source-based explanation, followed by an explicit transfer limit.
     u['task']='Self-check: explain the main distinction on this page, then name an assumption you would verify in the opening case.'
     u['answer']=' '.join(technical[:2])+' Apply this only after checking the operating conditions and dependencies in the case; a source description is not evidence that the proposed system has passed a test.'
     u['noAnswer']=False
   u['timebox']='';u['notes']='Teaching focus: '+clean(u['title'])+'. '+u.get('notes','')
  if n==15:
   approaches=[]
   for sl in slides:
    if sl['slide'] in [12,13,14]:
     for table in sl['tables']:
      approaches += table[1:]
   by['X02A']['blocks']=[dict(t='raw',html='<div class="lp-tablewrap"><table class="lp-table"><thead><tr><th>Reuse approach</th><th>Purpose in the source</th></tr></thead><tbody>'+''.join('<tr><td>'+esc(row[0])+'</td><td>'+esc(row[1])+'</td></tr>' for row in approaches)+'</tbody></table></div>')]
   by['X02A']['sub']='The supplied source lists fifteen approaches. Use the planning factors to select a suitable approach.'
   by['X02A']['answer']='Frameworks provide a reusable application structure; libraries provide callable abstractions; product lines generalize an application family; configurable and integrated systems reuse larger applications. Select by fit, lifetime, team capability, criticality and platform, rather than by reuse alone.'
  if n==11:
   by['X04']['blocks'].append(dict(t='wide',c='sand',lab='UNITS MATTER',txt='ROCOF may use time or transaction exposure in the supplied source (slide 71). State the denominator: failures per transaction are not failures per hour. The reciprocal is in the same exposure unit, under a compatible stable-rate interpretation. A sample ratio is a point estimate, not proof of the required population reliability.',src='P1 · pp.19–22, 71'))
  if n==16:
   by['D16']['blocks'].append(dict(t='raw',html='<div class="lp-callout"><b>Source example: addItem</b><pre style="white-space:pre-wrap;font:17px/1.7 monospace">context addItem\npre: PhotoLibrary.libSize() > 0\n     PhotoLibrary.retrieve(pid) = null\npost: libSize() = libSize()@pre + 1\n      PhotoLibrary.retrieve(pid) = p\n      PhotoLibrary.catEntry(pid) = photodesc</pre><p>The source assumes an existing library. @pre means the value before the call. Source slides 54–55.</p></div>'))
  # Opening context, outcomes and a visible pacing contract replace multiple orientation slides.
  case=by['R01'];goals=[]
  for b in by.get('R02R03',{}).get('blocks',[]):
   if b.get('t')=='rows':goals=[clean(x.get('txt','')) for x in b.get('items',[])][:5]
  goals=c['exit']
  case_text=case.get('sub','');decision=next((i.get('txt','') for b in case.get('blocks',[]) for i in b.get('items',[]) if 'DECISION' in i.get('lab','')),'')
  sessions='two teaching blocks of about 60–75 minutes' if c['weeks']==2 else 'one teaching block of about 75–90 minutes'
  intro=f'<div class="lp-callout"><p><b>Today’s artifact:</b> {esc(c["artifact"])}.</p><p><b>Fictional teaching case:</b> {case_text}</p><p>{decision}</p></div><h2 style="font:700 25px/1.4 system-ui">Five chapter learning objectives</h2><ol class="lp-outcomes">'+''.join('<li>'+g+'</li>' for g in goals)+'</ol>'+f'<p style="font:17px/1.6 system-ui"><b>Pacing suggestion:</b> {sessions}, plus 20–30 minutes of targeted source review. These are planning estimates, not an official timetable. Pause after Check 2 if teaching in two blocks. If time runs out, finish the remaining source topics in the next block; they stay required.</p><div class="lp-actions"><button class="lp-btn" data-jump="C01">See the source-to-lesson map</button></div>'
  intro=intro.replace('plus 20–30 minutes of targeted source review','plus the named required reading and a five-question preparation check (initial planning estimate: 15–25 minutes total; report your actual time)')
  intro+='<div class="lp-actions"><button class="lp-btn" data-jump="READING">Required self-study: exact pages and task</button><button class="lp-btn" data-jump="PREP">Five-objective preparation check</button></div><p class="lp-study-note">Complete this preparation before starting the chapter assignment and before the next chapter discussion. Blackboard supplies the calendar date. The five objectives, classroom concepts and named self-study are in chapter assessment scope; the optional toolkit is enrichment unless explicitly assigned. Do not repeat a completed submission under this new edition unless instructed.</p>'
  by['START']=rawunit('START','Your route through this chapter',intro);by['TITLE'].update(route='core',appendix=False)
  for idx,(q,a) in enumerate(c['checks'],1):by['CHECK'+str(idx)]=dict(k='CHECK'+str(idx),title='Check '+str(idx)+' · recall before revealing',phase='CHECK',route='core',blocks=[dict(t='wide',c='sand',lab='60-SECOND PAUSE',txt=q)],task='Think individually for 30 seconds, compare with a partner, then reveal the explanation.',answer=a,timebox='60 sec',appendix=False)
  # One concise decision practice keeps the methodology in the teaching route.
  by['APPLY']=dict(k='APPLY',title='Apply it · defend one bounded decision',phase='APPLY',route='core',appendix=False,
   blocks=[dict(t='wide',c='sand',lab='TEACHING CASE',txt=case_text),dict(t='cards',cols=3,items=[dict(c='teal',lab='CLAIM + EVIDENCE',txt='State a proposed action. Name the observed fact or source concept that supports it. Do not invent a completed test.',field=dict(id='path_claim',rows=3,ph='Proposed action and supporting evidence…')),dict(c='sand',lab='LIMIT + CHALLENGE',txt='Name an assumption and an observation that would reverse the decision.',field=dict(id='path_limit',rows=3,ph='The claim holds only if… It changes when…')),dict(c='green',lab='OWNER + NEXT STEP',txt='Name who checks the evidence and what must happen next.',field=dict(id='path_owner',rows=3,ph='The reviewer and the next verification step…'))]),dict(t='raw',html='<div class="lp-actions"><button class="lp-btn" data-jump="R07">Full six-field Decision Card</button><button class="lp-btn" data-jump="R15">AI verification gate</button></div><p style="font:16px/1.6 system-ui">If AI helped, name its contribution, check the claims yourself and retain human responsibility. This is practice; the separate assignment records your response to new evidence.</p>')],task='Work for four minutes, explain to a partner, then revise one sentence. The instructor samples reasoning, not the number of filled boxes.',answer=case.get('answer','')+' This is one defensible approach, not the only acceptable decision. Separate facts from requested evidence and name the claim’s limit.',timebox='6 min')
  for field in ['path_claim','path_limit','path_owner']:L.setdefault('fieldNames',{})['f:'+field]=field.replace('path_','').upper()
  L.setdefault('caseMap',[]).append(dict(title='Classroom application',parts=[['Claim and evidence','f:path_claim'],['Limit and challenge','f:path_limit'],['Owner and next step','f:path_owner']]))
  corekeys=['TITLE','START'];parts=[c['core'][:4],c['core'][4:8],c['core'][8:]]
  for i,chunk in enumerate(parts):
   corekeys+=chunk
   if i<2:corekeys+=['CHECK'+str(i+1)]
  corekeys+=['APPLY','END'];by['END'].update(route='core',appendix=False,phase='CHECK',notes='Finish with self-check, required source review and the after-class assignment.',task='')
  # Full source text/tables/embedded source pictures are preserved, not merely declared covered.
  studykeys=[];rows=[]
  bounds=[('Opening and source outline',1,2,['START'])]+[(g[0],g[1],min(g[2],count),g[3]) for g in c['groups']]
  covered=set()
  for gi,(title,a,b,keys) in enumerate(bounds):
   group_sl=[s for s in slides if a<=s['slide']<=b];covered|={s['slide'] for s in group_sl};key=f'SRC{gi:02d}'
   body='<div class="lp-callout"><p>Original supplied source slides '+str(a)+'–'+str(b)+'. The source retains historical examples and wording. Use the classroom explanation for contextual qualifications. Source review is part of the chapter.</p><p><a href="'+L['path']['sourceFile']+'" download>Download the original source pack</a> for full slide layout and vector diagrams.</p></div>'+''.join(source_html(s) for s in group_sl)
   by[key]=rawunit(key,title,body,route='study',sourceRange=f'{a}–{b}');studykeys.append(key)
   links=' · '.join('<button class="lp-btn" data-jump="'+k+'">'+esc(clean(by[k].get('title',k)))+'</button>' for k in keys if k in by)
   rows.append('<tr><td>'+esc(title)+'</td><td>'+str(a)+'–'+str(b)+'</td><td>'+links+'</td><td><button class="lp-btn" data-jump="'+key+'">Read source</button></td></tr>')
  if covered!=set(range(1,count+1)):raise ValueError(f'Source coverage incomplete: {n}: '+str(set(range(1,count+1))-covered))
  by['C01']=rawunit('C01','Source coverage and study checklist',f'<div class="lp-callout"><p><b>{count} supplied slides indexed.</b> Every original topic has a classroom location and a source review location. This map records material availability, not time actually taught or student mastery.</p><p>Five chapter objectives connect to supplied syllabus CLO '+', '.join(map(str,c['clo']))+'. Original concepts remain required even when detailed examples are completed after class.</p></div><div class="lp-tablewrap"><table class="lp-table"><thead><tr><th>Original topic</th><th>Source slides</th><th>Classroom location</th><th>Full source</th></tr></thead><tbody>'+''.join(rows)+'</tbody></table></div>',route='study')
  reading=study_mod.STUDY[n];required='<div class="lp-callout"><p><b>Required self-study · before the assignment and next chapter discussion.</b> Read the two selections below, attempt the five-question check, and use one assigned source slide to support the assignment artifact. Blackboard provides the calendar deadline.</p><p><b>Time budget:</b> initially 15–25 minutes including the check; record actual time. These replace unfocused source browsing. The original source ledger remains available for clarification and unfinished classroom topics.</p></div>'
  for title,a,b in reading['readings']:
   required+='<div class="lp-reading-row"><h3>'+esc(title)+'</h3><p>Original source slides '+str(a)+'–'+str(b)+'</p><div class="lp-actions">'+''.join('<button class="lp-btn" data-source-slide="'+str(p)+'">Slide '+str(p)+'</button>' for p in range(a,b+1))+'</div></div>'
  required+='<p class="lp-study-note"><b>Observable work:</b> identify the source slide and explain how its concept changes or supports your artifact. A citation or a checked box alone is insufficient. In the next class, be ready to explain a related example in your own words.</p><p class="lp-study-note"><b>Assessment scope:</b> all five chapter objectives, the classroom route and these named selections. Full source examples support this scope; archived administrative/title slides and optional toolkit activities are not extra assessed requirements. The instructor must announce any further assigned source topic in Blackboard before assessing it.</p><div class="lp-actions"><button class="lp-btn" data-jump="PREP">Attempt the five-objective check</button><button class="lp-btn" data-jump="C01">Full source map</button></div>'
  by['READING']=rawunit('READING','Required self-study · pages, purpose and deadline',required,route='study')
  quiz='<p class="lp-study-note">One question per chapter objective, in the same order as the five objectives. Attempt all five before opening feedback. Retry after reviewing an error. This public practice is not a secure test; the instructor uses a changed example during class.</p><form id="prepQuiz" class="lp-prep">'
  for qi,(q,options,answer,why) in enumerate(reading['quiz']):
   quiz+='<fieldset><legend>LO '+str(qi+1)+' · '+esc(q)+'</legend>'+''.join('<label><input type="radio" name="prep-'+str(qi)+'" value="'+str(oi)+'">'+esc(opt)+'</label>' for oi,opt in enumerate(options))+'<p hidden data-prep-feedback id="prep-feedback-'+str(qi)+'" class="prep-feedback"></p></fieldset>'
  quiz+='<button type="button" class="lp-btn" id="checkPrep">Check my five answers</button><div id="prepStatus" role="status" aria-live="polite">Attempt first. Feedback is initially hidden.</div></form>'
  by['PREP']=rawunit('PREP','Check all five objectives',quiz,route='study')
  ordered=[by[k] for k in corekeys]+[by['READING'],by['PREP'],by['C01']]+[by[k] for k in studykeys]+[u for k,u in by.items() if k not in corekeys+studykeys+['C01','READING','PREP']]
  L['units']=ordered;L['mainUnits']=corekeys;L['appendixUnits']=[u['k'] for u in ordered if u['route']!='core']
  for u in ordered:u['appendix']=u['route']!='core'
  # Remove copied administrative claims from all retained toolkit text.
  data=json.dumps(L,ensure_ascii=False).replace('before Assignment 1','before the chapter assignment').replace('this is what Assignment 1 asks for','reusable practice artifact').replace('CLO1','LO1').replace('CLO2','LO2').replace('CLO3','LO3').replace('CLO4','LO4').replace('CLO5','LO5')
  data=data.replace('</script','<\\/script')
  engine=engine.replace('Chapter 10 source content','this chapter’s source content').replace('MODEL ANSWER / INSTRUCTOR KEY','EXAMPLE EXPLANATION').replace("var KEY='iscarb-'+LECTURE.id", "var KEY='iscarb-'+LECTURE.id+'-path-v20260921'")
  # Safely convert the inherited JavaScript unit data into the reviewed source of truth.
  scripts=soup.find_all('script');scripts[0].string=prefix+'var LECTURE='+data+';\n'+engine
  for s in scripts[1:]:s.decompose()
  soup.body['class']='learning-path';soup.html['data-iscarb-standalone']='1';soup.html['data-learning-path']=BUILD
  style=soup.new_tag('style');style['id']='learning-path-style';style.string=css;soup.head.append(style)
  run=soup.new_tag('script');run.string=runtime;soup.body.append(run)
  for tag in soup.find_all(['script','link']):
   if 'kaspersky' in str(tag.get('src',''))+str(tag.get('href','')):tag.decompose()
  final=str(soup);path=f'lectures/iscarb/Ch{n}-{c["slug"]}.html';(ROOT/path).write_text(final)
  images=re.findall(r'data:image/(?:png|jpe?g|webp);base64,([A-Za-z0-9+/=]+)',final)
  origimgs=re.findall(r'data:image/(?:png|jpe?g|webp);base64,([A-Za-z0-9+/=]+)',markup)
  missing=set(origimgs)-set(images)
  if missing:raise ValueError('An original embedded figure was removed')
  record=dict(chapter=n,path=path,source_filename=src_html.name,source_sha256=hashlib.sha256(final.encode()).hexdigest(),original_html_sha256=hashlib.sha256(src_html.read_bytes()).hexdigest(),original_source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),source_download=destsource,embedded_image_sha256=[hashlib.sha256(base64.b64decode(im)).hexdigest() for im in images],source_slide_count=count,classroom_units=len(corekeys),original_embedded_images_preserved=len(set(origimgs)),core_keys=corekeys,source_groups=[dict(title=t,start=a,end=b,units=keys) for t,a,b,keys in bounds])
  catalog.append(record);audit.append({k:v for k,v in record.items() if k!='embedded_image_sha256'})
  print(f'CH{n}: {len(corekeys)} classroom units, {count} indexed source slides, {len(set(origimgs))} original embedded images preserved; {len(final)//1024} KB')
 (ROOT/'curriculum/learning-path/lectures.json').write_text(json.dumps(catalog,indent=2))
 (ROOT/'curriculum/learning-path/source-audit.json').write_text(json.dumps(audit,indent=2))
if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--input',type=Path,required=True);args=ap.parse_args();build(args.input)
